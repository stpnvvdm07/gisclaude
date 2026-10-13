# deck.json -> deck.pptx: редактируемая презентация (PowerPoint / Keynote / Google Slides / импорт в Figma).
# Текст — текстовые блоки (шрифт SF Pro Display), плашки и линии — фигуры, знак севера и значки легенд — кривые,
# карты, мазки, фото и логотип — картинки. Места под фото — серые фигуры с именем «ФОТО: …» (заменить через «Заменить рисунок»).
# Запуск: python3 export_pptx.py
import os, json, math, re
from PIL import Image
from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.oxml.ns import qn
from lxml import etree

H = os.path.dirname(os.path.abspath(__file__))
PX = 6350                                    # EMU в пикселе слайда 1920×1080 (13,333 × 7,5 дюйма)
E = lambda v: Emu(int(round(v * PX)))
FONT = "SF Pro Display"
def rgb(h): h = h.lstrip("#"); return RGBColor(int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16))
def set_alpha_fill(shape, op):
    if op >= 0.999: return
    sf = shape.fill._xPr.find(qn("a:solidFill"))
    if sf is not None and len(sf):
        a = etree.SubElement(sf[0], qn("a:alpha")); a.set("val", str(int(op * 100000)))
def no_line(shape): shape.line.fill.background()

def add_rect(sl, e, kind=MSO_SHAPE.RECTANGLE):
    sh = sl.shapes.add_shape(kind, E(e["x"]), E(e["y"]), E(e["w"]), E(e["h"]))
    f = e.get("fill")
    if f and f.startswith("#"): sh.fill.solid(); sh.fill.fore_color.rgb = rgb(f); set_alpha_fill(sh, e.get("op", 1))
    else: sh.fill.background()
    if e.get("stroke"):
        sh.line.color.rgb = rgb(e["stroke"]); sh.line.width = Pt(e.get("sw", 1) * 0.5)
        if e.get("dash"): sh.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    else: no_line(sh)
    sh.shadow.inherit = False
    if e.get("name"): sh.name = e["name"]
    return sh

def add_text(sl, e):
    lines = e["txt"].count("\n") + 1; h = max(e.get("h", e["size"] * e["lh"]), e["size"] * e["lh"] * lines)
    w = e["w"]
    if e.get("rot"):
        W = e.get("tw", w) + 10; cx, cy = e["x"] + h / 2, e["y"] - W / 2
        tb = sl.shapes.add_textbox(E(cx - W / 2), E(cy - h / 2), E(W), E(h)); tb.rotation = 270
    else:
        tb = sl.shapes.add_textbox(E(e["x"]), E(e["y"]), E(w + 2), E(h))
    tf = tb.text_frame; tf.word_wrap = not e.get("rot"); tf.auto_size = None
    for m in ("margin_left", "margin_right", "margin_top", "margin_bottom"): setattr(tf, m, 0)
    tf.vertical_anchor = MSO_ANCHOR.TOP
    for i, line in enumerate(e["txt"].split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = {"left": PP_ALIGN.LEFT, "right": PP_ALIGN.RIGHT, "center": PP_ALIGN.CENTER}[e["align"]]
        p.line_spacing = Pt(e["size"] * e["lh"] * 0.5)
        r = p.add_run(); r.text = line
        f = r.font; f.size = Pt(e["size"] * 0.5); f.name = FONT; f.color.rgb = rgb(e["color"])
        f.bold = e["wt"] in ("Semibold", "Bold")
        if e["wt"] == "Medium": f.name = FONT + " Medium"
        if e.get("op", 1) < 0.999:
            sf = r._r.get_or_add_rPr().find(qn("a:solidFill"))
            if sf is not None: a = etree.SubElement(sf[0], qn("a:alpha")); a.set("val", str(int(e["op"] * 100000)))
    return tb

_crop_cache = {}
def add_img(sl, e):
    src = os.path.join(H, e["src"])
    if not os.path.exists(src): return None
    im = Image.open(src); iw, ih = im.size; bw, bh = e["w"], e["h"]
    pic = sl.shapes.add_picture(src, E(e["x"]), E(e["y"]), E(bw), E(bh))
    fit = e.get("fit", "cover")
    if fit in ("cover", "contain") and iw and ih:
        ri, rb = iw / ih, bw / bh
        if fit == "cover":
            if ri > rb: c = (1 - rb / ri) / 2; pic.crop_left = pic.crop_right = c
            elif ri < rb: c = (1 - ri / rb) / 2; pic.crop_top = pic.crop_bottom = c
        else:
            if ri > rb: nh = bw / ri; pic.top = E(e["y"] + (bh - nh) / 2); pic.height = E(nh)
            elif ri < rb: nw = bh * ri; pic.left = E(e["x"] + (bw - nw) / 2); pic.width = E(nw)
    if e.get("circle"): pic.auto_shape_type = MSO_SHAPE.OVAL
    if e.get("op", 1) < 0.999:
        blip = pic._element.find(".//" + qn("a:blip")); a = etree.SubElement(blip, qn("a:alphaModFix")); a.set("amt", str(int(e["op"] * 100000)))
    return pic

def path_contours(d):
    toks = re.findall(r"[MLCZ]|-?\d+\.?\d*", d); i = 0; cur = None; out = []; pt = (0, 0)
    while i < len(toks):
        t = toks[i]
        if t == "M": pt = (float(toks[i + 1]), float(toks[i + 2])); cur = [pt]; out.append(cur); i += 3
        elif t == "L": pt = (float(toks[i + 1]), float(toks[i + 2])); cur.append(pt); i += 3
        elif t == "C":
            p1, p2, p3 = [(float(toks[i + 1 + 2 * k]), float(toks[i + 2 + 2 * k])) for k in range(3)]
            for s in range(1, 9):
                u = s / 8; v = 1 - u
                cur.append((v ** 3 * pt[0] + 3 * v * v * u * p1[0] + 3 * v * u * u * p2[0] + u ** 3 * p3[0],
                            v ** 3 * pt[1] + 3 * v * v * u * p1[1] + 3 * v * u * u * p2[1] + u ** 3 * p3[1]))
            pt = p3; i += 7
        elif t == "Z": i += 1
        else: i += 1
    return [c for c in out if len(c) >= 3]
def add_path(sl, e):
    cs = path_contours(e["d"])
    if not cs: return None
    ox, oy = e["x"], e["y"]
    fb = sl.shapes.build_freeform(E(ox + cs[0][0][0]), E(oy + cs[0][0][1]), scale=1.0)
    for k, c in enumerate(cs):
        if k: fb.move_to(E(ox + c[0][0]), E(oy + c[0][1]))
        fb.add_line_segments([(E(ox + x), E(oy + y)) for x, y in c[1:]], close=True)
    sh = fb.convert_to_shape()
    if e.get("fill"): sh.fill.solid(); sh.fill.fore_color.rgb = rgb(e["fill"]); set_alpha_fill(sh, e.get("op", 1))
    else: sh.fill.background()
    if e.get("stroke"): sh.line.color.rgb = rgb(e["stroke"]); sh.line.width = Pt(e.get("sw", 1) * 0.5)
    else: no_line(sh)
    return sh
def add_line(sl, e):
    c = sl.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, E(e["x1"]), E(e["y1"]), E(e["x2"]), E(e["y2"]))
    c.line.color.rgb = rgb(e["stroke"]); c.line.width = Pt(e["sw"] * 0.5)
    if e.get("dash"): c.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    return c
def add_arc(sl, e):
    n = 40; a0, a1 = e["a0"], e["a1"]; cx, cy, r, ir = e["cx"], e["cy"], e["r"], e["ir"]
    outer = [(cx + r * math.cos(a0 + (a1 - a0) * k / n), cy + r * math.sin(a0 + (a1 - a0) * k / n)) for k in range(n + 1)]
    inner = [(cx + ir * math.cos(a1 - (a1 - a0) * k / n), cy + ir * math.sin(a1 - (a1 - a0) * k / n)) for k in range(n + 1)] if ir > 0 else [(cx, cy)]
    pts = outer + inner
    fb = sl.shapes.build_freeform(E(pts[0][0]), E(pts[0][1]), scale=1.0); fb.add_line_segments([(E(x), E(y)) for x, y in pts[1:]], close=True)
    sh = fb.convert_to_shape(); sh.fill.solid(); sh.fill.fore_color.rgb = rgb(e["fill"]); no_line(sh); return sh

def main():
    deck = json.load(open(os.path.join(H, "deck.json"), encoding="utf-8"))
    prs = Presentation(); prs.slide_width = E(1920); prs.slide_height = E(1080)
    blank = prs.slide_layouts[6]
    for s in deck:
        sl = prs.slides.add_slide(blank)
        bg = sl.background.fill; bg.solid(); bg.fore_color.rgb = rgb(s["bg"])
        for e in s["els"]:
            t = e["t"]
            try:
                if t == "rect": add_rect(sl, e)
                elif t == "ellipse": add_rect(sl, e, MSO_SHAPE.OVAL)
                elif t == "text": add_text(sl, e)
                elif t == "img": add_img(sl, e)
                elif t == "path": add_path(sl, e)
                elif t == "line": add_line(sl, e)
                elif t == "arc": add_arc(sl, e)
            except Exception as ex:
                print("пропущен элемент", t, ex)
        if s.get("name"): sl.notes_slide.notes_text_frame.text = s["name"]
    out = os.path.join(H, "deck.pptx"); prs.save(out)
    print("deck.pptx:", len(deck), "слайдов,", round(os.path.getsize(out) / 1e6, 1), "МБ")
if __name__ == "__main__": main()
