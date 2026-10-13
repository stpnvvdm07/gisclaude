# Таблица xlsx v13: сводный счёт, показатели, чувствительность, устойчивость, допущения, освоенность, контроль x освоенность, изменения.
import os, warnings
import pandas as pd, geopandas as gpd
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
warnings.filterwarnings("ignore")
HERE = os.path.dirname(os.path.abspath(__file__)); V = os.path.join(HERE, "..", "vector")
R = pd.read_pickle(os.path.join(HERE, "v13_results.pkl"))
OUT = os.path.join(HERE, "..", "tables", "Сводный счёт нарушений, освоенность, чувствительность v13.xlsx")
vc = gpd.read_file(os.path.join(V, "problems_v13.gpkg"), layer="viol_cells")
osv = gpd.read_file(os.path.join(V, "analiz_v13.gpkg"), layer="osvoennost"); oc = gpd.read_file(os.path.join(V, "analiz_v13.gpkg"), layer="osv_cells")
S = vc[vc.in_scope == 1].sort_values("num")
FL = ["p1_1", "p1_2", "p1_3", "p1_4", "p2_1", "p2_2", "p2_3", "p3_1", "p3_2", "p3_3", "p3_4"]
svod = S[["num", "is_anal", "comfort_score"] + FL + ["p3_4a", "p3_4b", "p3_4c", "score", "cls", "score_wo13", "cls_wo13"]].copy()
svod.columns = ["Ячейка", "В зафикс. границах", "Балл комфорта"] + [c.replace("p", "").replace("_", ".") for c in FL] + \
               ["3.4а конкуренция", "3.4б неконтр.", "3.4в неосв.", "Счёт", "Класс", "Счёт без 1.3", "Класс без 1.3"]
total = {c: svod[c].sum() for c in svod.columns[3:14]}; total.update({"Ячейка": "Всего"})
svod = pd.concat([svod, pd.DataFrame([total])], ignore_index=True)
pok = S[["num", "pod", "nodes_non", "wear40", "loading_m2", "danger_obj", "combined_m", "transit_res_m", "biz_transit", "pg_out", "comp_pct", "unc_pct", "unosv_pct", "fav_closed"]].copy()
pok.columns = ["Ячейка", "Подъездов в 8 м", "Неорг. узлов", "Зданий износ ≥40 %", "Загрузка, м²", "Опасных объектов", "Совмещение, м",
               "Высокий транзит в АЖ/ПЖ, м", "Обсл. объектов у транзита", "Площадок вне контроля", "Конкуренция, %", "Неконтролируемая, %", "Неосвоенная, %", "Благопр. закрытых"]
ass = osv[osv.cls != "не оценивается"]
o1 = osv.groupby("cls").agg(Площадь_м2=("area", "sum"), Участков=("area", "size")).round(0).reset_index().rename(columns={"cls": "Класс"})
bu = ass.assign(B=ass.B.map({1: "да", 0: "нет"}), U=ass.U.map({1: "да", 0: "нет"})).groupby(["B", "U"]).area.sum().round(0).reset_index()
bu.columns = ["Благоустроено (B)", "Используется (U)", "Площадь, м²"]
uw = ass.groupby("U_why").area.sum().round(0).sort_values(ascending=False).reset_index(); uw.columns = ["Основание U", "Площадь, м²"]
bw = ass.groupby("B_why").area.sum().round(0).sort_values(ascending=False).reset_index(); bw.columns = ["Основание B", "Площадь, м²"]
nw = osv[osv.cls == "не оценивается"].groupby("note").area.sum().round(0).reset_index(); nw.columns = ["Не оценивается, причина", "Площадь, м²"]
ocs = oc[oc.in_scope == 1][["num", "assessed_m2", "osv_m2", "part_m2", "unosv_m2", "osv_pct", "part_pct", "unosv_pct", "parking_m2", "loading_m2"]].sort_values("num")
ocs.columns = ["Ячейка", "Оценено, м²", "Освоенная, м²", "Частично, м²", "Неосвоенная, м²", "Освоенная, %", "Частично, %", "Неосвоенная, %", "Стоянки, м²", "Загрузка, м²"]
metod = pd.DataFrame([
    ("Набор ячеек", "is_anal = 1 или подъезд в 8 м: 42 ячейки (34 + 14, 15, 22, 28, 29, 39, 47, 51)."),
    ("1.1", "Балл комфорта < 0,67: (вертикаль + замкнутость + размеры, по 0–2) / 6. Вертикаль 1,3–2,0 высокий; 0,6–1,3 средний. Замкнутость ≥ 0,65 / ≤ 0,35. Размер 25–140 м."),
    ("1.2", "В ячейку открывается ≥ 1 внешний неорганизованный узел (входы ближе 12 м; узел организован, если организованы все входы). Узел относится к ячейкам в зафиксированных границах в 8 м, иначе к ячейкам анализа в 8 м."),
    ("1.3", "Ячейка вне зафиксированных границ (is_anal = 0); в счёте 8 ячеек анализа. Дан вариант счёта без 1.3."),
    ("1.4", "Здание с износом ≥ 40 % в 8 м от ячейки (опубликованный износ; при отсутствии — середина оценки команды)."),
    ("2.1", "Зоны загрузки в ячейке ≥ 50 м² (по геометрии)."),
    ("2.2", "Опасный открытый объект в 6 м (бизнес «опасный + открытый»; мусорные площадки — допущение)."),
    ("2.3", "Участки «по проезжей части» с высокой интенсивностью пешеходного или транспортного дерева ≥ 20 м в ячейке."),
    ("3.1", "Ячейки с подъездами: высокий транзит в зонах АЖ/ПЖ ≥ 20 м."),
    ("3.2", "Обслуживающий объект в 6 м от ячейки и в 15 м от высокого пешеходного транзита."),
    ("3.3", "Место поведения ближней зоны вне АЖ; периферийной — вне ПЖ/ПП."),
    ("3.4", "а) конкуренция АЖ ∩ АП ≥ 5 % площади; б) неконтролируемая ≥ 50 %; в) неосвоенная ≥ 25 % оцениваемой площади. Флаг — любое из трёх."),
    ("Счёт", "Число пунктов из 11: 0–3 низкий, 4–5 средний, ≥ 6 высокий."),
    ("Освоенность v13", "Нарезка: площадки, деревья, трава (зелень DWG v9), прочее; ячейки, наблюдения, контур команды, учреждения. Не оцениваются: вне 42 ячеек, учреждения, «прочее» уже 5 м (после эрозии 2,5 м остаётся < 30 %). Площади по геометрии."),
], columns=["Пункт", "Правило"])
izm = pd.DataFrame([
    ("Зелень", "DWG v9 GREEN_AREAS (189 408 м²) вместо green v12 (171 939 м²); деревья по кронам v12 + массивам съёмки + лесу OSM."),
    ("Стоянки", "DWG v9: 21 (было 15), 6 новых — вид «открытая» (допущение)."),
    ("Здания", "make_valid 11 зданий с самопересечениями (26 и 8 — существенно)."),
    ("Ячейки", "Классы размеров 7 (139 м → комфортная) и 32 (142 м → некомфортная); вертикаль ячейки 1 — «низкий»; площади и коэффициенты по геометрии."),
    ("Узлы → ячейки", "Строго 8 м: 1.2 — 16 ячеек (было 19)."),
    ("Площади", "Все площади по геометрии (в v12 поле area освоенности и loading_m2/comp_pct завышены ~5 %)."),
    ("Освоенность", "Пересобрана на новой зелени и стоянках; критерий «уже 5 м» задан явно."),
    ("3.4", "Подпункты а/б/в в слоях и таблице."),
    ("Прочее", "Дубль фасада, тип 2 ограждений, bld_id бизнеса 55с3."),
], columns=["Что", "Изменение"])
sheets = [("Сводный счёт", svod), ("Показатели", pok), ("Чувствительность", R["SENS"]), ("Устойчивость ячеек", R["STAB"]), ("Допущения", R["ASS"]),
          ("Освоенность v13", None), ("Освоенность по ячейкам", ocs), ("Контроль × освоенность", R["XT"]), ("Методика счёта", metod), ("Изменения v13", izm)]
with pd.ExcelWriter(OUT, engine="openpyxl") as w:
    for name, df in sheets:
        if df is None:
            r = 0
            for part in (o1, bu, bw, uw, nw):
                part.to_excel(w, sheet_name=name, index=False, startrow=r); r += len(part) + 2
        else: df.to_excel(w, sheet_name=name, index=False)
    for ws in w.book.worksheets:
        for cell in ws[1]: cell.font = Font(bold=True, color="FFFFFF"); cell.fill = PatternFill("solid", fgColor="16161C")
        for col in ws.columns:
            L = max(len(str(c.value)) if c.value is not None else 0 for c in col); ws.column_dimensions[get_column_letter(col[0].column)].width = min(max(10, L + 2), 90)
            for c in col: c.alignment = Alignment(wrap_text=True, vertical="top")
print("->", OUT)
